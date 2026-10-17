#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Baut das Freelancer-Starter-Kit (DOCX, XLSX, PDF, CSV) und packt es als ZIP.

Aufruf:
    python3 business/product/build.py            # Vorschau-Build, Platzhalter erlaubt
    python3 business/product/build.py --strict   # Verkaufs-Build, bricht bei fehlender Konfiguration ab
"""
import csv
import json
import re
import shutil
import subprocess
import sys
import zipfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401  (Reserve)
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_COLOR_INDEX, WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE / "src"))

import mails  # noqa: E402
import onboarding  # noqa: E402
import prompts  # noqa: E402

CONFIG = json.loads((HERE.parent / "config.json").read_text(encoding="utf-8"))
STRICT = "--strict" in sys.argv

VERSION = "1.0"
PRODUCT = CONFIG.get("PRODUCT_NAME") or "Freelancer-Starter-Kit"
BUILD = HERE / "build"
DIST = HERE / "dist"

TEAL = RGBColor(0x0F, 0x4C, 0x5C)
GREY = RGBColor(0x55, 0x5B, 0x63)
FILL_BOX = "F1F5F7"
FILL_HEAD = "0F4C5C"
FILL_NOTE = "FFF6DD"

missing = []


def cfg(key, label):
    value = (CONFIG.get(key) or "").strip()
    if not value:
        missing.append(key)
        return f"[{label} EINTRAGEN]"
    return value


# ---------------------------------------------------------------- DOCX-Helfer

def new_doc(title):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)

    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.12

    for name, size, before, after in (("Heading 1", 18, 14, 6), ("Heading 2", 13, 12, 4), ("Heading 3", 11, 8, 2)):
        h = doc.styles[name]
        h.font.name = "Calibri"
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        h.element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        h.element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = TEAL
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.keep_with_next = True

    doc.core_properties.title = title
    doc.core_properties.author = CONFIG.get("SELLER_NAME") or PRODUCT
    add_footer(sec, f"{PRODUCT} · {title} · Seite ")
    return doc


def _field(run, code):
    for kind, text in (("begin", None), (None, code), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def add_footer(section, text):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.size = Pt(8.5)
    r.font.color.rgb = GREY
    r2 = p.add_run()
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = GREY
    _field(r2, "PAGE")


PLACEHOLDER = re.compile(r"(\[[^\]\n]+\])")


def rich(par, text, bold=False, italic=False, size=None, color=None, highlight=True):
    """Fügt Text ein und markiert [PLATZHALTER] gelb."""
    for part in PLACEHOLDER.split(text):
        if not part:
            continue
        run = par.add_run(part)
        run.bold = bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
        if highlight and PLACEHOLDER.fullmatch(part):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return par


def para(doc_or_cell, text="", bold=False, italic=False, size=None, color=None, align=None,
         after=None, highlight=True, style=None, keep_next=False):
    p = doc_or_cell.add_paragraph(style=style) if style else doc_or_cell.add_paragraph()
    if keep_next:
        p.paragraph_format.keep_with_next = True
    if text:
        rich(p, text, bold, italic, size, color, highlight)
    if align:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def cell_borders(cell, color="D5DDE2", size=4):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{k}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trPr.append(el)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    trPr.append(el)


def set_widths(table, widths_cm):
    table.autofit = False
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)


def box(doc, lines, fill=FILL_BOX, size=10.5, bold_first=False):
    """Einzelner umrahmter Textkasten; lines = Liste von Absätzen."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    no_split(t.rows[0])
    shade(c, fill)
    cell_borders(c)
    cell_margins(c, 110, 110, 160, 160)
    c.paragraphs[0].paragraph_format.space_after = Pt(0)
    first = True
    for i, line in enumerate(lines):
        p = c.paragraphs[0] if first else c.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(4)
        rich(p, line, bold=(bold_first and i == 0), size=size)
    set_widths(t, [16.6])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def table(doc, header, rows, widths, header_fill=FILL_HEAD, size=10, zebra=False):
    t = doc.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        shade(c, header_fill)
        cell_borders(c)
        cell_margins(c)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    repeat_header(t.rows[0])
    for n, row in enumerate(rows):
        cells = t.add_row().cells
        no_split(t.rows[-1])
        for i, val in enumerate(row):
            cell_borders(cells[i])
            cell_margins(cells[i])
            if zebra and n % 2:
                shade(cells[i], FILL_BOX)
            lines = val.split("\n") if isinstance(val, str) else [str(val)]
            for j, line in enumerate(lines):
                p = cells[i].paragraphs[0] if j == 0 else cells[i].add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                rich(p, line, size=size)
    set_widths(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def title_block(doc, title, subtitle):
    p = doc.add_paragraph()
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(26)
    r.font.color.rgb = TEAL
    p.paragraph_format.space_after = Pt(2)
    p2 = doc.add_paragraph()
    r2 = p2.add_run(subtitle)
    r2.font.size = Pt(12)
    r2.font.color.rgb = GREY
    p2.paragraph_format.space_after = Pt(14)


# ---------------------------------------------------------------- Dateien

def build_angebot(path):
    doc = new_doc("Angebotsvorlage")
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Normal"].paragraph_format.space_after = Pt(5)
    doc.sections[0].top_margin = doc.sections[0].bottom_margin = Cm(1.8)
    title_block(doc, "Angebotsvorlage", "Anleitung auf dieser Seite, Vorlage ab Seite 2. Anleitung vor dem Versand löschen.")
    doc.add_heading("So nutzt du die Vorlage", level=2)
    steps = [
        "Gelb markierte Stellen in eckigen Klammern [ ] sind Platzhalter. Ersetze sie und entferne die Markierung (Text markieren → Texthervorhebungsfarbe → „Keine Farbe“).",
        "Lösche alles, was für dieses Angebot nicht passt. Ein kurzes, klares Angebot wirkt professioneller als ein langes.",
        "Preise: Nutze die Tabelle „03_Kalkulation.xlsx“, um Stundensatz und Projektpreis herzuleiten, bevor du hier Beträge einträgst.",
        "Drei Pakete helfen Kunden bei der Entscheidung. Biete nur ein Paket an, wenn der Auftrag eindeutig ist. Dann lösche die Spalten der anderen Pakete.",
        "Umsatzsteuer: Wähle genau eine der beiden Varianten im Preisabschnitt (mit Umsatzsteuer oder Kleinunternehmerhinweis) und lösche die andere. Welche für dich gilt, klärst du mit deinem Steuerberater.",
        "Rechtliche Passagen (Haftung, Nutzungsrechte, Gerichtsstand, Widerruf, AGB) sind absichtlich nicht enthalten. Lass dafür eigene Geschäftsbedingungen oder einen Vertrag von einer Fachperson erstellen und hänge sie an.",
        "Pflichtangaben auf Rechnungen sind nicht Teil dieser Vorlage. Das ist ein Angebot, keine Rechnung.",
    ]
    for i, s in enumerate(steps, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(5)
        rich(p, f"{i}.  {s}")
    box(doc, ["Hinweis: Diese Vorlage ist keine Rechtsberatung. Prüfe Inhalt und Formulierungen auf deine Situation."],
        fill=FILL_NOTE, size=9.5)
    page_break(doc)

    # --- Kopf
    head = doc.add_table(rows=1, cols=2)
    left, right = head.rows[0].cells
    for c in (left, right):
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
    rich(left.paragraphs[0], "[DEIN NAME / FIRMA]", bold=True, size=11)
    for line in ("[Straße Nr.]", "[PLZ Ort]", "[E-Mail] · [Telefon]", "[Website]"):
        p = left.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        rich(p, line, size=9.5, color=GREY)
    rich(right.paragraphs[0], "Angebot Nr. [JJJJ-NNN]", bold=True, size=11)
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for line in ("Datum: [TT.MM.JJJJ]", "Gültig bis: [TT.MM.JJJJ]", "Ansprechperson: [NAME]"):
        p = right.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(0)
        rich(p, line, size=9.5, color=GREY)
    set_widths(head, [9.3, 7.3])
    para(doc, "", after=14)

    for line in ("[KUNDENFIRMA]", "z. Hd. [ANSPRECHPERSON]", "[Straße Nr.]", "[PLZ Ort]"):
        para(doc, line, after=0)
    para(doc, "", after=10)

    p = doc.add_paragraph()
    rich(p, "Angebot: [PROJEKTNAME]", bold=True, size=18, color=TEAL)
    para(doc, "Hallo [VORNAME],")
    para(doc, "vielen Dank für das Gespräch am [DATUM] und dein Vertrauen. Hier mein Vorschlag, wie wir [ZIEL IN EINEM SATZ] erreichen. Am Ende dieses Angebots findest du die nächsten Schritte.")

    doc.add_heading("1. Ausgangslage und Ziel", level=2)
    para(doc, "[2–4 Sätze: Was ist die Situation beim Kunden? Welches Problem soll gelöst werden? Woran erkennt ihr, dass das Projekt erfolgreich war?]")

    doc.add_heading("2. Leistungsumfang", level=2)
    para(doc, "Ich biete dir drei Pakete an. Das mittlere passt für die meisten Fälle.", after=4)
    table(
        doc,
        ["", "Basis", "Standard (empfohlen)", "Premium"],
        [
            ["Ergebnis", "[Kurzbeschreibung]", "[Kurzbeschreibung]", "[Kurzbeschreibung]"],
            ["Leistungen", "• [Leistung 1]\n• [Leistung 2]", "• Alles aus Basis\n• [Leistung 3]\n• [Leistung 4]", "• Alles aus Standard\n• [Leistung 5]\n• [Leistung 6]"],
            ["Korrekturschleifen", "[1]", "[2]", "[3]"],
            ["Lieferzeit", "[X Wochen]", "[X Wochen]", "[X Wochen]"],
            ["Preis (netto)", "[BETRAG] €", "[BETRAG] €", "[BETRAG] €"],
        ],
        [3.2, 4.4, 4.6, 4.4], size=9.5,
    )

    doc.add_heading("3. Nicht im Angebot enthalten", level=2)
    for s in ("[Beispiel: Erstellung von Bildmaterial oder Texten durch mich]",
              "[Beispiel: Laufende Betreuung nach Projektende]",
              "[Beispiel: Zusatzleistungen, die über die genannten Korrekturschleifen hinausgehen]"):
        para(doc, "• " + s, after=2)
    para(doc, "Wünschst du eine dieser Leistungen, erstelle ich dir gern ein Zusatzangebot.", after=4)

    doc.add_heading("4. Annahmen und Mitwirkung", level=2)
    for s in ("Du stellst bis [DATUM] folgende Unterlagen und Zugänge bereit: [LISTE].",
              "Feedback erfolgt innerhalb von [X] Werktagen. Verzögerungen auf deiner Seite verschieben den Zeitplan entsprechend.",
              "Es gibt eine feste Ansprechperson auf deiner Seite, die Entscheidungen trifft oder einholt.",
              "[WEITERE ANNAHME, ZUM BEISPIEL: Das Angebot beruht auf dem Stand unseres Gesprächs vom [DATUM].]"):
        para(doc, "• " + s, after=2)

    doc.add_heading("5. Zeitplan", level=2)
    table(
        doc,
        ["Meilenstein", "Ergebnis", "Termin"],
        [["Start", "Kickoff, Klärung offener Fragen", "[TT.MM.JJJJ]"],
         ["[Meilenstein 2]", "[Ergebnis]", "[TT.MM.JJJJ]"],
         ["[Meilenstein 3]", "[Ergebnis]", "[TT.MM.JJJJ]"],
         ["Abschluss", "Übergabe und Abnahme", "[TT.MM.JJJJ]"]],
        [4.2, 8.2, 4.2], size=9.5,
    )

    doc.add_heading("6. Preise und Zahlung", level=2)
    para(doc, "Gewähltes Paket: [Basis / Standard / Premium]", bold=True, after=4)
    table(
        doc,
        ["Position", "Betrag"],
        [["Projektpreis netto", "[BETRAG] €"],
         ["Umsatzsteuer [19] %", "[BETRAG] €"],
         ["Gesamtbetrag", "[BETRAG] €"]],
        [11.6, 5.0], size=10,
    )
    para(doc, "Variante für Kleinunternehmer (nur eine Variante behalten, Absprache mit Steuerberater): "
              "„Gemäß § 19 UStG wird keine Umsatzsteuer berechnet.“", italic=True, size=9.5, color=GREY)
    para(doc, "Zahlungsbedingungen (Variante wählen, andere löschen):", bold=True, after=2)
    para(doc, "A) [50] % bei Auftragserteilung, [50] % bei Lieferung. Zahlungsziel jeweils [14] Tage.", after=2)
    para(doc, "B) 100 % nach Abschluss des Projekts. Zahlungsziel [14] Tage.", after=2)
    para(doc, "C) Monatliche Abrechnung nach Aufwand zu [STUNDENSATZ] € netto pro Stunde, Obergrenze ohne Rücksprache: [BETRAG] €.", after=6)

    doc.add_heading("7. Gültigkeit und Annahme", level=2)
    para(doc, "Dieses Angebot gilt bis zum [TT.MM.JJJJ]. Zur Annahme genügt eine schriftliche Bestätigung per E-Mail oder die Unterschrift unten. [Es gelten meine beigefügten Geschäftsbedingungen / der beigefügte Vertrag.]")

    doc.add_heading("8. Nächste Schritte", level=2)
    for i, s in enumerate(("Du bestätigst das Angebot (Paket nennen).",
                           "Wir vereinbaren den Kickoff-Termin.",
                           "Ich starte am [TT.MM.JJJJ]."), 1):
        para(doc, f"{i}.  {s}", after=2)
    para(doc, "Ich freue mich auf die Zusammenarbeit!", keep_next=True)
    para(doc, "[DEIN NAME]", bold=True, after=10, keep_next=True)

    sig = doc.add_table(rows=2, cols=2)
    sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c in sig.rows[0].cells:
        c.paragraphs[0].paragraph_format.space_after = Pt(22)
        c.paragraphs[0].paragraph_format.keep_with_next = True
    rich(sig.rows[1].cells[0].paragraphs[0], "Ort, Datum", size=9, color=GREY)
    rich(sig.rows[1].cells[1].paragraphs[0], "Unterschrift Auftraggeber/in (Paket: [  ])", size=9, color=GREY)
    for c in sig.rows[1].cells:
        tcPr = c._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        top = OxmlElement("w:top")
        top.set(qn("w:val"), "single")
        top.set(qn("w:sz"), "6")
        top.set(qn("w:color"), "555B63")
        b.append(top)
        tcPr.append(b)
    set_widths(sig, [8.3, 8.3])
    doc.save(path)


def build_onboarding(path):
    doc = new_doc("Kunden-Onboarding-Fragebogen")
    title_block(doc, "Kunden-Onboarding-Fragebogen", f"{onboarding.count()} Fragen, mit denen du Projekte sauber startest")
    doc.add_heading("So nutzt du den Fragebogen", level=2)
    for s in ("Schicke den Fragebogen nach der Auftragsbestätigung, vor dem Kickoff. Bitte um Rückgabe innerhalb von [3–5] Werktagen.",
              "Streiche Fragen, die für das Projekt nicht passen. Weniger, dafür beantwortete Fragen sind besser als ein vollständiger, aber unbeantworteter Bogen.",
              "Nutze die Antworten für Kickoff, Angebotsabgleich und Abnahmekriterien. Unklare Antworten klärst du im Kickoff.",
              "Der Bogen nutzt die Sie-Form. Ersetze sie, wenn du deine Kunden duzt."):
        para(doc, "• " + s, after=2)
    para(doc, "Bitte den Hinweistext unten stehen lassen oder anpassen und diese Seite 1 vor dem Versand löschen.", italic=True, color=GREY, size=9.5)
    page_break(doc)

    p = doc.add_paragraph()
    rich(p, "Fragebogen zum Projektstart: [PROJEKTNAME]", bold=True, size=18, color=TEAL)
    para(doc, "Guten Tag [VORNAME / NAME],")
    para(doc, "damit das Projekt gut startet, bitte ich Sie, die folgenden Fragen zu beantworten. Stichworte genügen. Wenn Sie etwas nicht wissen, schreiben Sie „offen“, dann klären wir es gemeinsam im Kickoff. Bitte senden Sie den ausgefüllten Bogen bis [DATUM] zurück an [E-MAIL].")
    para(doc, "Projekt: [PROJEKTNAME]   ·   Datum: [TT.MM.JJJJ]   ·   Ausgefüllt von: [NAME]", size=9.5, color=GREY)

    n = 0
    for section, qs in onboarding.SECTIONS:
        doc.add_heading(section, level=2)
        t = doc.add_table(rows=0, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for q in qs:
            n += 1
            cells = t.add_row().cells
            no_split(t.rows[-1])
            for c in cells:
                cell_borders(c)
                cell_margins(c, 90, 90, 120, 120)
            shade(cells[0], FILL_BOX)
            cells[0].paragraphs[0].paragraph_format.space_after = Pt(0)
            rich(cells[0].paragraphs[0], f"{n}. {q}", size=10)
            cells[1].paragraphs[0].paragraph_format.space_after = Pt(26)
        set_widths(t, [8.3, 8.3])
    para(doc, "", after=6)
    para(doc, "Vielen Dank! Ich melde mich nach Eingang innerhalb von [X] Werktagen mit dem Terminvorschlag für den Kickoff.")
    doc.save(path)


def build_mails(path):
    doc = new_doc("Mail-Textbausteine")
    title_block(doc, "Mail-Textbausteine", f"{len(mails.MAILS)} Vorlagen für Akquise, Angebot, Projekt, Geld und Kundenbeziehung")
    doc.add_heading("So nutzt du die Bausteine", level=2)
    for s in ("Kopiere den Text in dein Mailprogramm. Gelb markierte Stellen in eckigen Klammern [ ] ersetzt du.",
              "Kürze, was nicht passt. Eine Mail soll nach dir klingen. Wenn du siezt, ändere Anrede und Formulierungen.",
              "Schicke keine Mail, die du nicht gelesen hast. Prüfe Betrag, Datum, Namen und Anhänge.",
              "Hinweise in Klammern mit „Hinweis für dich“ gehören nicht in die Mail. Lösche sie vor dem Versand."):
        para(doc, "• " + s, after=2)

    doc.add_heading("Inhalt", level=2)
    n = 0
    for g in mails.groups():
        items = [m for m in mails.MAILS if m[0] == g]
        para(doc, f"{g} ({len(items)})", bold=True, after=1)
        for m in items:
            n += 1
            para(doc, f"    {n}. {m[1]}", size=10, after=0)
    page_break(doc)

    n = 0
    for g in mails.groups():
        doc.add_heading(g, level=1)
        for grp, title, when, subject, body in [m for m in mails.MAILS if m[0] == g]:
            n += 1
            doc.add_heading(f"{n}. {title}", level=3)
            para(doc, "Wann: " + when, italic=True, size=9.5, color=GREY, after=3, keep_next=True)
            t = doc.add_table(rows=1, cols=1)
            c = t.rows[0].cells[0]
            no_split(t.rows[0])
            shade(c, FILL_BOX)
            cell_borders(c)
            cell_margins(c, 110, 110, 160, 160)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(6)
            rich(p, "Betreff: ", bold=True, size=10)
            rich(p, subject, size=10)
            for chunk in body.split("\n"):
                q = c.add_paragraph()
                q.paragraph_format.space_after = Pt(0)
                rich(q, chunk, size=10)
            set_widths(t, [16.6])
            para(doc, "", after=4)
    doc.save(path)


def build_prompts(path, csv_path):
    all_p = prompts.all_prompts()
    doc = new_doc("KI-Prompts")
    title_block(doc, "KI-Prompt-Bibliothek", f"{len(all_p)} Prompts für Selbstständige in {len(prompts.CATEGORIES)} Kategorien")
    doc.add_heading("So arbeitest du mit den Prompts", level=2)
    for s in (
        "Kopiere den Prompt in deinen KI-Assistenten (zum Beispiel ChatGPT, Claude oder Gemini). Ersetze die gelb markierten Platzhalter [ ] durch deine Angaben. Je konkreter, desto besser das Ergebnis.",
        "Gib so viel Kontext wie nötig, aber keine vertraulichen Kundendaten, Passwörter oder Gesundheitsdaten. Prüfe vorab, was dein Vertrag und der Datenschutz erlauben. Im Zweifel anonymisieren (siehe Prompt „Vertrauliche Daten anonymisieren“).",
        "Lies jedes Ergebnis kritisch. KI kann falsche Fakten, Zahlen und Quellen erfinden. Du bist verantwortlich für alles, was du an Kunden schickst.",
        "Steuer-, Rechts- und Finanzthemen: Die Prompts helfen dir beim Vorbereiten und Strukturieren, nicht beim Entscheiden. Lass verbindliche Fragen von Fachleuten beantworten.",
        "Füge am Ende eines Prompts gern hinzu: „Stelle mir Rückfragen, bevor du antwortest, wenn dir etwas fehlt.“ Das verbessert Ergebnisse spürbar.",
        "Speichere Prompts, die bei dir gut funktionieren. Die CSV-Datei „05_KI-Prompts.csv“ lässt sich in Google Sheets, Excel oder Notion importieren.",
    ):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(4)
        rich(p, "•  " + s)

    doc.add_heading("Inhalt", level=2)
    start = 1
    for cat, desc, items in prompts.CATEGORIES:
        para(doc, f"{cat}  ·  Nr. {start}–{start + len(items) - 1}", bold=True, after=0)
        para(doc, "    " + desc, size=9.5, color=GREY, after=3)
        start += len(items)
    page_break(doc)

    n = 0
    for cat, desc, items in prompts.CATEGORIES:
        doc.add_heading(cat, level=1)
        para(doc, desc, italic=True, color=GREY, after=6, keep_next=True)
        for title, text in items:
            n += 1
            h = doc.add_heading(f"{n}. {title}", level=3)
            h.paragraph_format.keep_with_next = True
            t = doc.add_table(rows=1, cols=1)
            c = t.rows[0].cells[0]
            no_split(t.rows[0])
            shade(c, FILL_BOX)
            cell_borders(c)
            cell_margins(c, 100, 100, 150, 150)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            rich(p, text, size=10)
            set_widths(t, [16.6])
            para(doc, "", after=2)
    doc.save(path)

    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Nr", "Kategorie", "Titel", "Prompt"])
        for row in all_p:
            w.writerow(row)


# ---------------------------------------------------------------- XLSX

def build_xlsx(path):
    wb = Workbook()
    blue = Font(name="Calibri", size=11, color="0000FF")
    bold = Font(name="Calibri", size=11, bold=True)
    head_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=FILL_HEAD)
    input_fill = PatternFill("solid", fgColor="EAF3FF")
    result_fill = PatternFill("solid", fgColor="E6F4EA")
    note_fill = PatternFill("solid", fgColor=FILL_NOTE)
    thin = Side(style="thin", color="C9D3D8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    eur = '#,##0.00 "€"'
    eur0 = '#,##0 "€"'
    pct = "0%"

    def title(ws, text, sub):
        ws["A1"] = text
        ws["A1"].font = Font(name="Calibri", size=18, bold=True, color="0F4C5C")
        ws["A2"] = sub
        ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="555B63")

    # --- Anleitung
    ws = wb.active
    ws.title = "Anleitung"
    title(ws, "Kalkulation für Selbstständige", f"{PRODUCT} · Version {VERSION}")
    lines = [
        ("So funktioniert die Tabelle", True),
        ("1. Blatt „Stundensatz“: Trage deine Wünsche und Kosten in die blauen Felder ein. Das Ergebnis zeigt, welchen Stundensatz du mindestens brauchst.", False),
        ("2. Blatt „Projektkalkulation“: Zerlege ein Projekt in Arbeitspakete. Die Tabelle rechnet Puffer, Preis und drei Pakete aus und prüft, ob dein Zielstundensatz erreicht wird.", False),
        ("3. Blatt „Auslastung“: Plane 12 Wochen im Voraus, wie viele Stunden gebucht sind und wo Lücken oder Überlastung entstehen.", False),
        ("", False),
        ("Farben", True),
        ("Blaue Schrift auf hellblauem Hintergrund = Eingabefeld. Grüne Felder = Ergebnisse. Alles andere sind Formeln. Bitte nicht überschreiben.", False),
        ("", False),
        ("Wichtig", True),
        ("Die vorausgefüllten Werte sind Beispiele und keine Empfehlung. Passe sie an deine Situation an.", False),
        ("Die Tabelle ist ein Rechenhilfsmittel und keine Steuer- oder Rechtsberatung. Abgabenquote, Umsatzsteuer und Kosten klärst du mit deinem Steuerberater.", False),
        ("Die Beispiele in „Projektkalkulation“ (Arbeitspakete, Stunden) sind frei erfunden und dienen nur dem Verständnis.", False),
    ]
    for i, (text, is_head) in enumerate(lines, start=4):
        c = ws.cell(row=i, column=1, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name="Calibri", size=12 if is_head else 11, bold=is_head, color="0F4C5C" if is_head else "000000")
    ws.column_dimensions["A"].width = 110
    for r in range(4, 4 + len(lines)):
        ws.row_dimensions[r].height = 32 if len(lines[r - 4][0]) > 100 else 18

    # --- Stundensatz
    ws = wb.create_sheet("Stundensatz")
    title(ws, "Stundensatz berechnen", "Blaue Felder ausfüllen · Beispielwerte überschreiben")
    inputs = [
        (4, "Wunsch-Netto pro Monat (€)", 3000, eur0, "Was soll nach Steuern und Vorsorge privat verfügbar sein?"),
        (5, "Abgabenquote auf den Gewinn (%)", 0.30, pct, "Schätzung für Steuern, Kranken-, Pflege- und Rentenvorsorge. Mit Steuerberater abstimmen."),
        (6, "Geschäftskosten pro Monat (€)", 400, eur0, "Software, Versicherungen, Büro, Weiterbildung, Marketing …"),
        (7, "Puffer / Rücklage auf den Umsatzbedarf (%)", 0.10, pct, "Für Flauten, Anschaffungen und Unvorhergesehenes."),
        (8, "Arbeitstage pro Woche", 5, "0", ""),
        (9, "Arbeitsstunden pro Tag", 8, "0.0", ""),
        (10, "Urlaubstage pro Jahr", 25, "0", "Tage, die auf deine Arbeitstage fallen."),
        (11, "Feiertage pro Jahr (auf Arbeitstage)", 10, "0", "Je nach Bundesland und Arbeitsmuster."),
        (12, "Krankheits- und Ausfalltage pro Jahr", 8, "0", "Realistisch einplanen."),
        (13, "Anteil abrechenbarer Stunden (%)", 0.60, pct, "Akquise, Verwaltung und Weiterbildung sind nicht abrechenbar. Häufig 50 bis 70 Prozent."),
        (14, "Kleinunternehmerregelung (§ 19 UStG)?", "Nein", None, "Ja oder Nein. Bei Unsicherheit: Steuerberater fragen."),
        (15, "Umsatzsteuersatz (%)", 0.19, pct, "Nur relevant bei „Nein“ oben."),
    ]
    ws["A3"], ws["B3"], ws["C3"] = "Eingabe", "Wert", "Hinweis"
    for c in ("A3", "B3", "C3"):
        ws[c].font, ws[c].fill = head_font, head_fill
    for row, label, val, fmt, hint in inputs:
        ws.cell(row=row, column=1, value=label).border = border
        c = ws.cell(row=row, column=2, value=val)
        c.font, c.fill, c.border = blue, input_fill, border
        if fmt:
            c.number_format = fmt
        h = ws.cell(row=row, column=3, value=hint)
        h.font = Font(name="Calibri", size=10, italic=True, color="555B63")
    dv = DataValidation(type="list", formula1='"Ja,Nein"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B14")

    ws["A17"], ws["B17"], ws["C17"] = "Berechnung", "Ergebnis", "Rechenweg"
    for c in ("A17", "B17", "C17"):
        ws[c].font, ws[c].fill = head_font, head_fill
    calc = [
        (18, "Arbeitstage pro Jahr", "=MAX(0,52*B8-B10-B11-B12)", "0", "52 Wochen × Arbeitstage − Urlaub − Feiertage − Ausfall"),
        (19, "Verfügbare Stunden pro Jahr", "=B18*B9", "#,##0", "Arbeitstage × Stunden pro Tag"),
        (20, "Abrechenbare Stunden pro Jahr", "=B19*B13", "#,##0", "Verfügbare Stunden × abrechenbarer Anteil"),
        (21, "Benötigter Gewinn vor Abgaben (€/Jahr)", "=B4*12/(1-B5)", eur0, "Wunsch-Netto × 12 ÷ (1 − Abgabenquote)"),
        (22, "Geschäftskosten (€/Jahr)", "=B6*12", eur0, "Monatliche Kosten × 12"),
        (23, "Umsatzbedarf ohne Puffer (€/Jahr)", "=B21+B22", eur0, "Gewinn + Kosten"),
        (24, "Umsatzziel inklusive Puffer (€/Jahr)", "=B23*(1+B7)", eur0, "Umsatzbedarf × (1 + Puffer)"),
    ]
    for row, label, f, fmt, hint in calc:
        ws.cell(row=row, column=1, value=label).border = border
        c = ws.cell(row=row, column=2, value=f)
        c.number_format, c.border = fmt, border
        ws.cell(row=row, column=3, value=hint).font = Font(name="Calibri", size=10, italic=True, color="555B63")

    res = [
        (26, "Dein Mindest-Stundensatz (netto)", '=IF(B20>0,B24/B20,"Eingaben prüfen")', eur),
        (27, "Entspricht einem Tagessatz (netto)", '=IF(ISNUMBER(B26),B26*B9,"–")', eur),
        (28, "Stundensatz brutto (inkl. Umsatzsteuer, falls zutreffend)", '=IF(ISNUMBER(B26),IF(B14="Ja",B26,B26*(1+B15)),"–")', eur),
    ]
    for row, label, f, fmt in res:
        a = ws.cell(row=row, column=1, value=label)
        a.font, a.border = bold, border
        c = ws.cell(row=row, column=2, value=f)
        c.number_format, c.fill, c.border, c.font = fmt, result_fill, border, Font(name="Calibri", size=12, bold=True)

    ws["A30"] = "Wie sich der Satz bei anderer Auslastung ändert"
    ws["A30"].font = Font(name="Calibri", size=12, bold=True, color="0F4C5C")
    ws["A31"], ws["B31"], ws["C31"] = "Anteil abrechenbar", "Stundensatz netto", ""
    for c in ("A31", "B31"):
        ws[c].font, ws[c].fill = head_font, head_fill
    for i, q in enumerate((0.5, 0.6, 0.7, 0.8), start=32):
        a = ws.cell(row=i, column=1, value=q)
        a.number_format, a.border = pct, border
        a.alignment = Alignment(horizontal="left")
        b = ws.cell(row=i, column=2, value=f'=IF(B$19*A{i}>0,B$24/(B$19*A{i}),"–")')
        b.number_format, b.border = eur, border
    ws["A37"] = "Hinweis: Das ist eine Rechenhilfe. Sie bildet weder Steuerberechnung noch Marktpreise ab. Prüfe das Ergebnis gegen das, was dein Markt zahlt."
    ws["A37"].font = Font(name="Calibri", size=10, italic=True, color="555B63")
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 70

    # --- Projektkalkulation
    ws = wb.create_sheet("Projektkalkulation")
    title(ws, "Projektkalkulation", "Arbeitspakete schätzen, Puffer einrechnen, Pakete bilden · Beispielwerte überschreiben")
    ws["A4"], ws["B4"] = "Zielstundensatz (netto)", "=Stundensatz!B26"
    ws["C4"] = "Wird aus dem Blatt „Stundensatz“ übernommen. Du kannst hier auch einen festen Wert eintragen."
    ws["B4"].number_format = eur
    ws["B4"].fill, ws["B4"].border, ws["B4"].font = result_fill, border, Font(name="Calibri", size=12, bold=True)
    ws["C4"].font = Font(name="Calibri", size=10, italic=True, color="555B63")

    ws["A6"], ws["B6"] = "Unsicherheit", "Zuschlag auf Stunden"
    for c in ("A6", "B6"):
        ws[c].font, ws[c].fill = head_font, head_fill
    for i, (lvl, f) in enumerate((("niedrig", 1.10), ("mittel", 1.25), ("hoch", 1.50)), start=7):
        ws.cell(row=i, column=1, value=lvl).border = border
        c = ws.cell(row=i, column=2, value=f)
        c.font, c.fill, c.border = blue, input_fill, border
        c.number_format = '0.00"×"'
    ws["C7"] = "niedrig: du hast das schon oft gemacht · mittel: Erfahrung, aber Details offen · hoch: neu oder unklar"
    ws["C7"].font = Font(name="Calibri", size=10, italic=True, color="555B63")

    hdr = ["Arbeitspaket", "Geschätzte Stunden", "Unsicherheit", "Stunden inkl. Puffer", "Kosten netto"]
    for i, h in enumerate(hdr, start=1):
        c = ws.cell(row=11, column=i, value=h)
        c.font, c.fill, c.border = head_font, head_fill, border
        c.alignment = Alignment(wrap_text=True, vertical="center")
    sample = [("Beispiel: Briefing und Konzept", 6, "niedrig"),
              ("Beispiel: Erste Umsetzung", 18, "mittel"),
              ("Beispiel: Abstimmung und Korrekturen", 8, "hoch"),
              ("Beispiel: Finalisierung und Übergabe", 4, "niedrig")]
    first, last = 12, 26
    dv2 = DataValidation(type="list", formula1='"niedrig,mittel,hoch"', allow_blank=True)
    ws.add_data_validation(dv2)
    for r in range(first, last + 1):
        idx = r - first
        a = ws.cell(row=r, column=1, value=sample[idx][0] if idx < len(sample) else None)
        b = ws.cell(row=r, column=2, value=sample[idx][1] if idx < len(sample) else None)
        c = ws.cell(row=r, column=3, value=sample[idx][2] if idx < len(sample) else None)
        for x in (a, b, c):
            x.font, x.fill, x.border = blue, input_fill, border
        b.number_format = "0.0"
        dv2.add(c)
        d = ws.cell(row=r, column=4, value=f'=IF(B{r}="","",B{r}*IFERROR(VLOOKUP(C{r},$A$7:$B$9,2,FALSE),1))')
        e = ws.cell(row=r, column=5, value=f'=IF(D{r}="","",D{r}*$B$4)')
        d.number_format, e.number_format = "0.0", eur
        d.border = e.border = border

    t = last + 1  # Summenzeile
    ws.cell(row=t, column=1, value="Summe").font = bold
    ws.cell(row=t, column=2, value=f"=SUM(B{first}:B{last})").number_format = "0.0"
    ws.cell(row=t, column=4, value=f"=SUM(D{first}:D{last})").number_format = "0.0"
    ws.cell(row=t, column=5, value=f"=SUM(E{first}:E{last})").number_format = eur
    for col in range(1, 6):
        ws.cell(row=t, column=col).font = bold
        ws.cell(row=t, column=col).border = border

    r0 = t + 2  # 29
    rows = [
        (r0, "Fremdkosten (Lizenzen, Material, Reisen) netto", 0, eur, True, "Eigene Auslagen, die du weitergibst."),
        (r0 + 1, "Aufschlag auf Fremdkosten (%)", 0.0, pct, True, "Optional, für Verwaltungsaufwand."),
        (r0 + 2, "Projektpreis Standard (netto)", f"=E{t}+B{r0}*(1+B{r0+1})", eur, False, "Arbeitszeit inkl. Puffer + Fremdkosten"),
    ]
    for r, label, v, fmt, is_input, hint in rows:
        ws.cell(row=r, column=1, value=label).border = border
        c = ws.cell(row=r, column=2, value=v)
        c.number_format, c.border = fmt, border
        if is_input:
            c.font, c.fill = blue, input_fill
        else:
            c.font, c.fill = Font(name="Calibri", size=12, bold=True), result_fill
        ws.cell(row=r, column=3, value=hint).font = Font(name="Calibri", size=10, italic=True, color="555B63")
    std = f"B{r0+2}"

    p0 = r0 + 4  # Pakete
    ws.cell(row=p0, column=1, value="Pakete").font = Font(name="Calibri", size=12, bold=True, color="0F4C5C")
    for i, h in enumerate(("Paket", "Faktor auf Standard", "Preis netto", "Preis brutto"), start=1):
        c = ws.cell(row=p0 + 1, column=i, value=h)
        c.font, c.fill, c.border = head_font, head_fill, border
    for i, (name, f) in enumerate((("Basis", 0.70), ("Standard", 1.00), ("Premium", 1.35)), start=p0 + 2):
        ws.cell(row=i, column=1, value=name).border = border
        fc = ws.cell(row=i, column=2, value=f)
        fc.font, fc.fill, fc.border = blue, input_fill, border
        fc.number_format = '0.00"×"'
        pr = ws.cell(row=i, column=3, value=f"=ROUND({std}*B{i},0)")
        pr.number_format, pr.border = eur0, border
        br = ws.cell(row=i, column=4, value=f'=IF(Stundensatz!$B$14="Ja",C{i},C{i}*(1+Stundensatz!$B$15))')
        br.number_format, br.border = eur0, border
    ws.cell(row=p0 + 5, column=1, value="Der Faktor bestimmt, wie viel Leistung (und Preis) jedes Paket relativ zu Standard umfasst. Die Leistungen der Pakete legst du im Angebot fest.").font = Font(name="Calibri", size=10, italic=True, color="555B63")

    c0 = p0 + 7  # Check
    ws.cell(row=c0, column=1, value="Prüfung: Erreichst du deinen Zielstundensatz?").font = Font(name="Calibri", size=12, bold=True, color="0F4C5C")
    chk = [
        (c0 + 1, "Effektiver Stundensatz, wenn der Aufwand wie geschätzt anfällt", f'=IF(B{t}>0,(B{r0+2}-B{r0}*(1+B{r0+1}))/B{t},"–")'),
        (c0 + 2, "Effektiver Stundensatz, wenn der Puffer ausgeschöpft wird", f'=IF(D{t}>0,(B{r0+2}-B{r0}*(1+B{r0+1}))/D{t},"–")'),
        (c0 + 3, "Effektiver Stundensatz bei 25 % mehr Aufwand als Puffer-Schätzung", f'=IF(D{t}>0,(B{r0+2}-B{r0}*(1+B{r0+1}))/(D{t}*1.25),"–")'),
    ]
    for r, label, f in chk:
        ws.cell(row=r, column=1, value=label).border = border
        c = ws.cell(row=r, column=2, value=f)
        c.number_format, c.border, c.fill = eur, border, result_fill
    ws.cell(row=c0 + 4, column=1, value="Bewertung").font = bold
    ws.cell(row=c0 + 4, column=2, value=f'=IF(NOT(ISNUMBER(B{c0+2})),"–",IF(B{c0+2}>=B4*0.99,"OK: Ziel auch bei ausgeschöpftem Puffer erreicht","Achtung: Ziel wird bei vollem Aufwand verfehlt – Preis oder Umfang anpassen"))')
    ws.cell(row=c0 + 4, column=2).font = Font(name="Calibri", size=11, bold=True)
    ws.cell(row=c0 + 4, column=2).fill = note_fill
    green = PatternFill("solid", bgColor="C6EFCE")
    red = PatternFill("solid", bgColor="FFC7CE")
    ws.conditional_formatting.add(f"B{c0+1}:B{c0+3}", CellIsRule(operator="lessThan", formula=["$B$4*0.99"], fill=red))
    ws.conditional_formatting.add(f"B{c0+1}:B{c0+3}", CellIsRule(operator="greaterThanOrEqual", formula=["$B$4*0.99"], fill=green))

    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 16
    ws.row_dimensions[11].height = 32
    ws.freeze_panes = "A12"

    # --- Auslastung
    ws = wb.create_sheet("Auslastung")
    title(ws, "Auslastung – 12 Wochen", "Gebuchte Stunden eintragen · Zeile rot = überbucht, gelb = Luft nach oben")
    ws["A3"], ws["B3"] = "Verfügbare Stunden pro Woche", "=ROUND(Stundensatz!B9*Stundensatz!B8,0)"
    ws["B3"].number_format = "0"
    ws["B3"].font, ws["B3"].fill, ws["B3"].border = Font(name="Calibri", size=11, bold=True), result_fill, border
    ws["C3"] = "Standard aus dem Blatt „Stundensatz“. Du kannst pro Woche abweichen (Urlaub, Feiertage)."
    ws["C3"].font = Font(name="Calibri", size=10, italic=True, color="555B63")
    hdr = ["Woche", "Verfügbar (h)", "Kunde / Projekt A", "Kunde / Projekt B", "Kunde / Projekt C", "Akquise & Verwaltung",
           "Gebucht (h)", "Auslastung", "Frei (h)"]
    for i, h in enumerate(hdr, start=1):
        c = ws.cell(row=5, column=i, value=h)
        c.font, c.fill, c.border = head_font, head_fill, border
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[5].height = 34
    for w in range(12):
        r = 6 + w
        ws.cell(row=r, column=1, value=f"Woche {w+1}").border = border
        ws.cell(row=r, column=2, value="=$B$3").border = border
        for col in range(3, 7):
            c = ws.cell(row=r, column=col, value=None)
            c.font, c.fill, c.border = blue, input_fill, border
            c.number_format = "0.0"
        g = ws.cell(row=r, column=7, value=f"=SUM(C{r}:F{r})")
        g.number_format, g.border = "0.0", border
        a = ws.cell(row=r, column=8, value=f'=IF(B{r}>0,G{r}/B{r},"–")')
        a.number_format, a.border = pct, border
        f_ = ws.cell(row=r, column=9, value=f"=B{r}-G{r}")
        f_.number_format, f_.border = "0.0", border
    ws.conditional_formatting.add("H6:H17", CellIsRule(operator="greaterThan", formula=["1"], fill=red))
    ws.conditional_formatting.add("H6:H17", CellIsRule(operator="lessThan", formula=["0.6"], fill=PatternFill("solid", bgColor="FFEB9C")))
    ws["A19"] = "Summe / Durchschnitt"
    ws["B19"] = "=SUM(B6:B17)"
    for col, letter in zip(range(3, 8), "CDEFG"):
        ws.cell(row=19, column=col, value=f"=SUM({letter}6:{letter}17)")
    ws["H19"] = '=IF(B19>0,G19/B19,"–")'
    ws["H19"].number_format = pct
    ws["I19"] = "=B19-G19"
    for c in range(1, 10):
        ws.cell(row=19, column=c).font = bold
        ws.cell(row=19, column=c).border = border
    ws.column_dimensions["A"].width = 24
    for letter in "BCDEFGHI":
        ws.column_dimensions[letter].width = 16
    ws.freeze_panes = "A6"

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
    wb.save(path)


# ---------------------------------------------------------------- START HIER

def build_start(path):
    seller = cfg("SELLER_NAME", "NAME DES ANBIETERS")
    email = cfg("CONTACT_EMAIL", "SUPPORT-E-MAIL")
    doc = new_doc("Start hier")
    title_block(doc, f"{PRODUCT}", f"Version {VERSION} · Stand {date.today().strftime('%m/%Y')} · Start hier")
    para(doc, "Willkommen und danke für deinen Kauf! Dieses Kit spart dir die Fleißarbeit bei Angeboten, Kalkulation und Kundenkommunikation. Alles ist auf Deutsch, alles ist anpassbar.")

    doc.add_heading("Das ist drin", level=2)
    table(
        doc,
        ["Datei", "Wofür", "Programm"],
        [
            ["01_Angebotsvorlage.docx", "Angebot mit drei Paketen, Zeitplan, Preis- und Zahlungsabschnitt, Annahmeblock", "Word, LibreOffice, Google Docs"],
            ["02_Kunden-Onboarding-Fragebogen.docx", f"{onboarding.count()} Fragen in 7 Abschnitten für einen sauberen Projektstart", "Word, LibreOffice, Google Docs"],
            ["03_Kalkulation.xlsx", "Stundensatz-Rechner, Projektkalkulation mit Puffer und Paketpreisen, 12-Wochen-Auslastung", "Excel, LibreOffice, Google Sheets"],
            ["04_Mail-Textbausteine.docx", f"{len(mails.MAILS)} Mail-Vorlagen: Akquise, Angebot, Projekt, Geld, Kundenbeziehung", "Word, LibreOffice, Google Docs"],
            ["05_KI-Prompts.docx", f"{len(prompts.all_prompts())} Prompts in {len(prompts.CATEGORIES)} Kategorien, bequem zum Kopieren", "Word, LibreOffice, Google Docs"],
            ["05_KI-Prompts.csv", "Dieselben Prompts zum Import in Notion, Google Sheets, Excel", "Tabellenkalkulation, Notion"],
        ],
        [5.6, 7.6, 3.4], size=9.5, zebra=True,
    )

    doc.add_heading("In 30 Minuten startklar", level=2)
    for i, s in enumerate((
        "Öffne „03_Kalkulation.xlsx“, Blatt „Stundensatz“, und trage deine Zahlen ein. Notiere dir dein Ergebnis (5 Min.).",
        "Öffne „01_Angebotsvorlage.docx“. Ersetze deine Absenderdaten und speichere die Datei als deine Standardvorlage (10 Min.).",
        "Passe „02_Kunden-Onboarding-Fragebogen.docx“ an deine Leistung an: streichen, was nicht passt (10 Min.).",
        "Wähle aus „04_Mail-Textbausteine.docx“ die drei Mails, die du am häufigsten brauchst, und lege sie in deinem Mailprogramm als Textbausteine ab (5 Min.).",
        "Wenn du mit KI arbeitest: Starte mit den Prompts „Stilprofil erstellen“ und „Vertrauliche Daten anonymisieren“ aus Kategorie 10.",
    ), 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(4)
        rich(p, f"{i}.  {s}")

    doc.add_heading("Wie du die Vorlagen nutzt", level=2)
    para(doc, "Gelb markierte Stellen in eckigen Klammern [ ] sind Platzhalter. Ersetze sie durch deine Angaben. Streiche, was für dich nicht passt. Kunden merken, wenn Vorlagen unverändert verschickt werden.")

    doc.add_heading("Nutzungsrechte", level=2)
    para(doc, "Du darfst alle Inhalte für dein eigenes Geschäft und deine Kundenprojekte nutzen und anpassen. Nicht erlaubt sind das Weiterverkaufen, das Verschenken, das Veröffentlichen zum Download und das Weitergeben an Dritte, auch nicht in veränderter Form als eigenes Produkt. Die Lizenz gilt für eine Person.")

    doc.add_heading("Wichtige Hinweise", level=2)
    for s in (
        "Die Inhalte sind keine Rechts-, Steuer- oder Anlageberatung. Verbindliche Fragen zu Verträgen, AGB, Umsatzsteuer, Rechnungsstellung und Datenschutz klärst du mit Fachleuten (Anwalt, Steuerberater).",
        "Beispielwerte in den Dateien (Preise, Quoten, Stunden) sind Beispiele und keine Empfehlung.",
        "KI-Ergebnisse können fehlerhaft sein. Prüfe alles, bevor du es verwendest oder verschickst, und gib keine vertraulichen Daten ein, die du nicht weitergeben darfst.",
        "Es gibt keine Garantie auf bestimmte Ergebnisse, Aufträge oder Einnahmen.",
    ):
        para(doc, "• " + s, after=2)

    doc.add_heading("Fragen oder Probleme?", level=2)
    para(doc, f"Schreib mir an {email}. Nenne bitte kurz, welche Datei und welches Programm du nutzt, dann komme ich schneller zu einer Antwort.")
    para(doc, f"© {date.today().year} {seller}. Alle Rechte vorbehalten.", size=9, color=GREY)
    doc.save(path)


# ---------------------------------------------------------------- Main

def to_pdf(docx_path, outdir):
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(docx_path)],
        check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180,
    )
    return outdir / (docx_path.stem + ".pdf")


def recalc_xlsx(path):
    """LibreOffice rechnet alle Formeln und speichert die Werte mit, damit Vorschauen
    (Handy, Mac-Vorschau, Mail-Clients) keine leeren Zellen zeigen."""
    tmp = path.parent / "_recalc"
    tmp.mkdir()
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp), str(path)],
        check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180,
    )
    shutil.move(str(tmp / path.name), str(path))
    tmp.rmdir()


def main():
    if BUILD.exists():
        shutil.rmtree(BUILD)
    BUILD.mkdir(parents=True)
    DIST.mkdir(exist_ok=True)
    root = BUILD / PRODUCT
    root.mkdir()

    build_angebot(root / "01_Angebotsvorlage.docx")
    build_onboarding(root / "02_Kunden-Onboarding-Fragebogen.docx")
    build_xlsx(root / "03_Kalkulation.xlsx")
    recalc_xlsx(root / "03_Kalkulation.xlsx")
    build_mails(root / "04_Mail-Textbausteine.docx")
    build_prompts(root / "05_KI-Prompts.docx", root / "05_KI-Prompts.csv")

    start_docx = BUILD / "00_START-HIER.docx"
    build_start(start_docx)
    pdf = to_pdf(start_docx, root)
    start_docx.unlink()

    if missing and STRICT:
        print("FEHLER: Konfiguration unvollständig (business/config.json):", ", ".join(sorted(set(missing))))
        sys.exit(1)

    suffix = "" if not missing else "_VORSCHAU"
    zip_path = DIST / f"{PRODUCT}_v{VERSION}{suffix}.zip"
    for old in DIST.glob(f"{PRODUCT}_v*.zip"):
        old.unlink()
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
        for f in sorted(root.iterdir()):
            z.write(f, f"{PRODUCT}/{f.name}")
    manifest = {
        "VERSION": VERSION,
        "PROMPT_COUNT": len(prompts.all_prompts()),
        "PROMPT_CATEGORIES": len(prompts.CATEGORIES),
        "MAIL_COUNT": len(mails.MAILS),
        "QUESTION_COUNT": onboarding.count(),
        "QUESTION_SECTIONS": len(onboarding.SECTIONS),
        "ZIP_KB": round(zip_path.stat().st_size / 1024),
    }
    (DIST / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print("OK:", zip_path.relative_to(HERE.parent.parent), f"({zip_path.stat().st_size/1024:.0f} KB)")
    if missing:
        print("Hinweis: Vorschau-Build. Fehlende Konfiguration:", ", ".join(sorted(set(missing))))
    return root


if __name__ == "__main__":
    main()
